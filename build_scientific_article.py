from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).parent / "docs" / "Avance_Articulo_Cientifico_Tienda_Piscinas.docx"
BLUE = "087FA8"
NAVY = "06364D"
LIGHT = "E8F8FA"
GRAY = "555555"

def font(run, size=10, bold=False, italic=False, color="000000", name="Times New Roman"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")

def set_table_widths(table, widths):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW"); tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total)); tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd"); tbl_ind.set(qn("w:w"), "100"); tbl_ind.set(qn("w:type"), "dxa"); tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[idx])); tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)

def set_columns(section, count=2, space=360):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    node = cols[0] if cols else OxmlElement("w:cols")
    node.set(qn("w:num"), str(count)); node.set(qn("w:space"), str(space)); node.set(qn("w:equalWidth"), "1")
    if not cols: sect_pr.append(node)

def add_para(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=3, first_line=True):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.0
    if first_line: p.paragraph_format.first_line_indent = Inches(.15)
    font(p.add_run(text), 10)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    font(p.add_run(text), 10 if level == 1 else 9.5, bold=True, color=NAVY)
    return p

def add_reference(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.18); p.paragraph_format.first_line_indent = Inches(-.18)
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
    font(p.add_run(f"[{number}] {text}"), 8)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5); section.page_height = Inches(11)
section.top_margin = Inches(.7); section.bottom_margin = Inches(.7)
section.left_margin = Inches(.72); section.right_margin = Inches(.72)
section.header_distance = Inches(.3); section.footer_distance = Inches(.35)

# Styles: standard_business_brief base with a named IEEE-draft override.
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"; normal.font.size = Pt(10)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.paragraph_format.space_after = Pt(3); normal.paragraph_format.line_spacing = 1.0
for name, before, after in (("Heading 1", 8, 3), ("Heading 2", 6, 2)):
    st = doc.styles[name]; st.font.name = "Times New Roman"; st.font.size = Pt(10 if name.endswith("1") else 9.5); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(NAVY)
    st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman"); st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after); st.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(header.add_run("AVANCE DE ARTÍCULO CIENTÍFICO · TIENDA PISCINAS SANTAMARÍA"), 7.5, color=GRAY, name="Arial")
footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Borrador académico · 17 de agosto de 2026"), 7.5, color=GRAY, name="Arial")

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
font(p.add_run("Diseño e implementación de una tienda web transaccional para la gestión de productos y servicios de piscinas"), 18, bold=True, color=NAVY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
font(p.add_run("Samuel [apellidos] · [Nombre del segundo integrante]"), 10.5)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
font(p.add_run("Desarrollo de Aplicaciones Web y Patrones · [Universidad] · Costa Rica\n[correo institucional 1] · [correo institucional 2]"), 9, italic=True, color=GRAY)

abstract = doc.add_paragraph(); abstract.paragraph_format.left_indent = Inches(.35); abstract.paragraph_format.right_indent = Inches(.35); abstract.paragraph_format.space_after = Pt(5)
font(abstract.add_run("Resumen—"), 9, bold=True)
font(abstract.add_run("Este trabajo presenta el diseño y el avance de implementación de Tienda Piscinas Santamaría, una aplicación web orientada a centralizar la consulta de productos, equipos y servicios para el mantenimiento de piscinas en Costa Rica. La solución adopta una arquitectura Modelo-Vista-Controlador sobre Java y Spring Boot, vistas dinámicas con Thymeleaf y persistencia relacional mediante Jakarta Persistence. El primer incremento implementado comprende una interfaz adaptable, un catálogo consultado desde la base de datos, un esquema reproducible de diez tablas y pruebas automatizadas del contexto y del recorrido MVC. La validación técnica ejecutó dos pruebas sin fallos y confirmó la recuperación y presentación de ocho productos de demostración. Los resultados son preliminares: autenticación, autorización, operaciones CRUD completas, internacionalización y cierre transaccional del pedido pertenecen al siguiente incremento. El aporte principal de esta fase es una base trazable y verificable que reduce el riesgo de integrar posteriormente seguridad y comercio electrónico."), 9)
kw = doc.add_paragraph(); kw.paragraph_format.left_indent = Inches(.35); kw.paragraph_format.right_indent = Inches(.35); kw.paragraph_format.space_after = Pt(10)
font(kw.add_run("Palabras clave—"), 9, bold=True); font(kw.add_run("aplicación web, Spring Boot, MVC, Thymeleaf, JPA, comercio electrónico, piscinas."), 9)

# Full-width snapshot table before the two-column body.
caption = doc.add_paragraph(); caption.alignment = WD_ALIGN_PARAGRAPH.CENTER; caption.paragraph_format.space_after = Pt(4)
font(caption.add_run("TABLA I. ESTADO VERIFICABLE DEL PRIMER INCREMENTO"), 8, bold=True, color=NAVY)
table = doc.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_widths(table, [2050, 4320, 1950])
header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
header_repeat = OxmlElement("w:tblHeader")
header_repeat.set(qn("w:val"), "true")
header_tr_pr.append(header_repeat)
for i, text in enumerate(("Componente", "Evidencia implementada", "Estado")):
    cell = table.rows[0].cells[i]; set_cell_shading(cell, NAVY); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; font(cell.paragraphs[0].add_run(text), 8, bold=True, color="FFFFFF", name="Arial")
rows = [
    ("Presentación", "Portada adaptable, navegación, categorías, productos, servicios y contacto.", "Implementado"),
    ("Persistencia", "10 tablas, relaciones, índices, catálogo JPA y datos de prueba.", "Implementado"),
    ("Transacción", "Pedido y detalle de demostración en el esquema.", "Diseño inicial"),
    ("Calidad", "2 pruebas automatizadas aprobadas; respuesta HTTP 200 verificada.", "Implementado"),
    ("Seguridad e i18n", "Autenticación, roles efectivos y selector de idioma.", "Siguiente fase"),
]
for ridx, row in enumerate(rows):
    cells = table.add_row().cells
    for i, text in enumerate(row):
        if ridx % 2: set_cell_shading(cells[i], LIGHT)
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
        font(cells[i].paragraphs[0].add_run(text), 8, name="Arial")

two = doc.add_section(WD_SECTION.CONTINUOUS)
two.top_margin = Inches(.7); two.bottom_margin = Inches(.7); two.left_margin = Inches(.72); two.right_margin = Inches(.72)
set_columns(two, 2, 360)

add_heading(doc, "I. INTRODUCCIÓN")
add_para(doc, "Las pequeñas empresas especializadas en mantenimiento de piscinas requieren canales digitales capaces de presentar inventario, orientar la selección de tratamientos y registrar solicitudes comerciales. Cuando la información permanece distribuida entre mensajes, llamadas y listados manuales, el cliente carece de una vista consistente de la oferta y el negocio pierde trazabilidad sobre productos, disponibilidad y transacciones. A partir de esta necesidad se planteó Tienda Piscinas Santamaría como una aplicación web que concentre productos químicos, equipos, accesorios y servicios técnicos.")
add_para(doc, "El problema de investigación aplicado se formula así: ¿cómo construir una base web mantenible y verificable que permita evolucionar desde un escaparate digital hacia una tienda transaccional con persistencia, seguridad y colaboración demostrable? La propuesta emplea Spring Boot para estructurar la aplicación y simplificar su ejecución [1]; Spring MVC para separar solicitudes, modelo y vistas [2]; Thymeleaf para renderizar HTML del lado del servidor [3]; y Jakarta Persistence para mapear objetos del dominio a una base relacional [4]. Bootstrap y estilos propios apoyan el comportamiento adaptable de la interfaz [5].")
add_para(doc, "El objetivo general es desarrollar una aplicación web transaccional para gestionar la oferta de Piscinas Santamaría y facilitar la interacción con clientes potenciales. En este corte, el objetivo específico es implementar y comprobar la base arquitectónica, visual y persistente correspondiente al primer 50 % del proyecto. El alcance incluye catálogo dinámico, esquema relacional, datos reproducibles, pruebas automatizadas y documentación; excluye temporalmente autenticación efectiva, CRUD administrativo completo, carrito definitivo, internacionalización y despliegue productivo.")
add_para(doc, "La relevancia académica reside en integrar contenidos del curso mediante un caso coherente: HTML5/CSS, Bootstrap, Spring Boot, MVC, Thymeleaf, JPA, SQL, pruebas y control de versiones. La relevancia práctica consiste en ofrecer un punto de partida que pueda validarse con un cliente potencial y extenderse sin reemplazar la arquitectura inicial.")

add_heading(doc, "II. METODOLOGÍA")
add_heading(doc, "A. Enfoque incremental y división del trabajo", 2)
add_para(doc, "Se adoptó un proceso incremental basado en ramas Git. El primer integrante trabajó en la rama samuel y produjo un corte vertical funcional: la solicitud HTTP llega al controlador, el repositorio recupera entidades desde la base de datos y Thymeleaf presenta los resultados. El uso de ramas permite aislar cambios y preparar revisiones mediante pull requests, práctica recomendada para el desarrollo colaborativo [10]. La segunda fase se asignó explícitamente al otro integrante para generar contribuciones distintas y auditables.")
add_para(doc, "Cada incremento sigue cuatro actividades: análisis del criterio de rúbrica, implementación, verificación automatizada y documentación. Los commits describen propósito y alcance; las decisiones técnicas registran alternativas, consecuencias y trabajo futuro. Este procedimiento facilita reconstruir la evolución del sistema durante la defensa.")
add_heading(doc, "B. Arquitectura de software", 2)
add_para(doc, "La solución utiliza Java 21 y Spring Boot 3.5.4. InicioController atiende la ruta raíz y solicita al ProductoRepository los productos activos. Las clases Producto y Categoria representan el dominio; la relación muchos-a-uno vincula cada producto con una categoría. Una anotación EntityGraph carga de manera explícita la categoría requerida por la vista. Esta decisión mantiene desactivado Open Session in View, evita consultas tardías durante el renderizado y delimita el acceso a datos conforme al ciclo de vida definido por Jakarta Persistence [4].")
add_para(doc, "La capa de presentación usa fragmentos Thymeleaf para encabezado y pie de página. La plantilla principal itera sobre la colección de productos y aplica expresiones para nombre, categoría, descripción y precio. Thymeleaf facilita esta integración con controladores y expresiones Spring EL [3]. Bootstrap 5.3 y una hoja de estilos propia implementan cuadrículas adaptables, navegación móvil, jerarquía visual y la paleta azul, celeste, blanco y turquesa [5].")
add_heading(doc, "C. Modelo y estrategia de datos", 2)
add_para(doc, "El script schema.sql crea diez tablas: roles, usuarios, usuarios_roles, clientes, categorias, productos, servicios, solicitudes_servicio, pedidos y detalles_pedido. Las claves foráneas conectan identidad, catálogo, solicitudes y compras. Pedidos funciona como cabecera transaccional, mientras detalles_pedido conserva producto, cantidad, precio unitario y subtotal. El modelo responde al estándar objeto-relacional de Jakarta Persistence [4] y prevé MySQL como motor productivo [7].")
add_para(doc, "Para asegurar reproducibilidad sin exigir infraestructura externa, el perfil de desarrollo utiliza H2 en memoria y ejecuta automáticamente schema.sql y data.sql [6]. Los datos iniciales incluyen cuatro categorías, ocho productos, tres servicios, dos roles, un cliente y un pedido de demostración. La contraseña de demostración se conserva deliberadamente inválida hasta que el segundo incremento incorpore BCrypt y Spring Security; por tanto, no se presenta como credencial funcional.")
add_heading(doc, "D. Verificación", 2)
add_para(doc, "La verificación se realizó con Maven y pruebas JUnit integradas en Spring Boot. La primera prueba comprueba que el contexto de aplicación se construya. La segunda ejecuta una solicitud simulada a la ruta raíz, exige estado satisfactorio, vista index, colección de ocho productos y presencia del producto Cloro granulado premium en el HTML generado. Además, se inició el servidor Tomcat embebido y se comprobó una respuesta HTTP 200 en localhost:8080.")

add_heading(doc, "III. RESULTADOS PRELIMINARES")
add_heading(doc, "A. Resultado funcional", 2)
add_para(doc, "El incremento produce una página navegable con propuesta de valor, indicadores demostrativos, categorías, catálogo, servicios, testimonios marcados como contenido provisional, formulario de contacto y acceso a WhatsApp. La interfaz se adapta mediante puntos de ruptura para escritorio, tableta y móvil. El catálogo no está codificado como contenido estático: cada tarjeta se genera a partir de entidades consultadas por el repositorio JPA, lo que demuestra un uso visible de la persistencia.")
add_para(doc, "La ejecución de la suite final reportó: Tests run: 2, Failures: 0, Errors: 0, Skipped: 0. Durante la primera ejecución, la prueba de integración reveló una LazyInitializationException al acceder a producto.categoria.nombre con la sesión cerrada. La corrección añadió una carga explícita mediante EntityGraph. Este hallazgo evidencia el valor de probar el recorrido completo y no solamente el arranque del contexto.")
add_heading(doc, "B. Resultado estructural", 2)
add_para(doc, "El esquema supera el mínimo de ocho tablas e incorpora una transacción representativa. Sin embargo, la existencia de tablas no equivale todavía a un flujo comercial completo. El resultado estratégico se alcanzará cuando una operación de negocio valide existencias, calcule importes en servidor, cree pedido y detalles en una transacción atómica y descuente inventario. Por ello, el presente artículo distingue entre estructura implementada y comportamiento pendiente.")
add_heading(doc, "C. Resultado de trazabilidad", 2)
add_para(doc, "El repositorio contiene README, instrucciones de ejecución, scripts SQL, informe del avance, lista de tareas del segundo integrante y registro de decisiones. La rama samuel conserva commits descriptivos y puede revisarse antes de integrar. Esta documentación conecta artefactos con criterios de almacenamiento, diseño, uso de tecnologías y colaboración.")

add_heading(doc, "IV. DISCUSIÓN")
add_para(doc, "Los resultados respaldan la viabilidad de una arquitectura Spring MVC para el dominio planteado. El recorrido persistente del catálogo disminuye el riesgo de que la base de datos sea decorativa, mientras el esquema relacional prepara relaciones y consultas posteriores. La decisión de mantener open-in-view desactivado hace visibles los límites de la capa de persistencia y obligó a definir la información necesaria en la consulta. Aunque esta política requiere mayor disciplina, evita que la vista active accesos implícitos difíciles de medir.")
add_para(doc, "H2 aporta reproducibilidad para pruebas y MySQL se reserva para el entorno final. Esta combinación presenta una amenaza de compatibilidad: pequeñas diferencias de tipos, funciones SQL o restricciones pueden aparecer al migrar. La mitigación propuesta es ejecutar las mismas pruebas con un perfil MySQL antes del despliegue y, en una versión posterior, adoptar migraciones versionadas.")
add_para(doc, "La interfaz sigue un enfoque responsive coherente con Bootstrap [5], pero no se afirma conformidad formal con WCAG. Las pautas WCAG 2.2 organizan la accesibilidad en los principios perceptible, operable, comprensible y robusto [8]. Una auditoría de contraste, foco, navegación por teclado, nombres accesibles y mensajes de error debe formar parte del segundo incremento.")
add_para(doc, "La principal limitación metodológica es la ausencia, hasta este corte, de evaluación con usuarios o evidencia directa del cliente. Las cifras de experiencia y los testimonios de la maqueta son contenido demostrativo y no resultados empíricos. Antes de la entrega deben sustituirse o respaldarse con una entrevista, encuesta o análisis de mercado. También están pendientes controles de autenticación y autorización; OWASP recomienda verificar estos mecanismos sistemáticamente antes de considerar segura una aplicación [9].")

add_heading(doc, "V. CONCLUSIONES PARCIALES Y TRABAJO FUTURO")
add_para(doc, "El primer incremento demuestra que Tienda Piscinas Santamaría puede consultar y presentar un catálogo persistente mediante una arquitectura MVC reproducible. Se implementaron una interfaz adaptable, diez tablas relacionadas, datos de demostración y dos pruebas automatizadas satisfactorias. La detección y corrección del problema de carga perezosa confirma que la estrategia de pruebas aporta evidencia técnica útil para la defensa.")
add_para(doc, "No se considera concluida la aplicación transaccional. El siguiente incremento debe implementar Spring Security con roles, CRUD de categorías y productos, solicitudes de servicio, carrito y creación atómica de pedidos, validación de inventario, internacionalización español/inglés, pruebas de autorización, configuración MySQL y despliegue. También debe recopilar evidencia de cliente potencial y ejecutar una evaluación básica de usabilidad y accesibilidad. Las conclusiones finales del artículo se actualizarán únicamente después de medir esas funciones.")

add_heading(doc, "REFERENCIAS")
refs = [
    "Spring, “Spring Boot 3.5 Reference Documentation,” 2026. [En línea]. Disponible: https://docs.spring.io/spring-boot/3.5/reference/index.html. [Accedido: 17-ago-2026].",
    "Spring, “Spring Web MVC,” Spring Framework Reference. [En línea]. Disponible: https://docs.spring.io/spring-framework/reference/web/webmvc.html. [Accedido: 17-ago-2026].",
    "The Thymeleaf Team, “Tutorial: Thymeleaf + Spring,” versión 3.1, 2026. [En línea]. Disponible: https://www.thymeleaf.org/doc/tutorials/3.1/thymeleafspring.html. [Accedido: 17-ago-2026].",
    "Eclipse Foundation, “Jakarta Persistence 3.1 Specification,” mar. 2022. [En línea]. Disponible: https://jakarta.ee/specifications/persistence/3.1/. [Accedido: 17-ago-2026].",
    "Bootstrap Team, “Get started with Bootstrap v5.3,” 2026. [En línea]. Disponible: https://getbootstrap.com/docs/5.3/getting-started/introduction/. [Accedido: 17-ago-2026].",
    "H2 Group, “H2 Database Engine Features.” [En línea]. Disponible: https://www.h2database.com/html/features.html. [Accedido: 17-ago-2026].",
    "Oracle, “MySQL 8.4 Reference Manual.” [En línea]. Disponible: https://dev.mysql.com/doc/refman/8.4/en/. [Accedido: 17-ago-2026].",
    "W3C Web Accessibility Initiative, “Web Content Accessibility Guidelines (WCAG) 2 Overview,” 2026. [En línea]. Disponible: https://www.w3.org/WAI/standards-guidelines/wcag/. [Accedido: 17-ago-2026].",
    "OWASP Foundation, “Application Security Verification Standard.” [En línea]. Disponible: https://owasp.org/www-project-application-security-verification-standard/. [Accedido: 17-ago-2026].",
    "GitHub, “Writing code for a project,” GitHub Docs, 2026. [En línea]. Disponible: https://docs.github.com/en/pull-requests/concepts/writing-code-for-a-project. [Accedido: 17-ago-2026].",
    "Spring, “Spring Data JPA Reference Documentation.” [En línea]. Disponible: https://docs.spring.io/spring-data/jpa/reference/. [Accedido: 17-ago-2026].",
]
for i, ref in enumerate(refs, 1): add_reference(doc, i, ref)

add_heading(doc, "NOTA DE CONTROL DEL BORRADOR")
add_para(doc, "Este documento constituye un avance basado exclusivamente en el primer 50 % implementado. Antes de la entrega final deben reemplazarse los campos entre corchetes, insertar capturas numeradas del sistema, incorporar evidencia del segundo integrante, actualizar resultados y discusión, y revisar cada referencia conforme a la plantilla IEEE suministrada por el docente.", first_line=False)

doc.core_properties.title = "Avance del artículo científico - Tienda Piscinas Santamaría"
doc.core_properties.subject = "Primer 50 % del proyecto de Desarrollo de Aplicaciones Web y Patrones"
doc.core_properties.author = "Equipo Tienda Piscinas Santamaría"
doc.core_properties.keywords = "Spring Boot, MVC, Thymeleaf, JPA, tienda web"
doc.save(OUT)
print(OUT)
